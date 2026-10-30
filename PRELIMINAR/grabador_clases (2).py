"""
Grabador de clases: graba -> transcribe (Whisper) -> resume (Gemini) -> exporta (Word / Google Docs)

Instalar:
    pip install sounddevice soundfile openai-whisper google-genai python-docx
    pip install google-api-python-client google-auth-oauthlib

En la Raspberry Pi solo necesitas la PARTE 1 (grabar).
Las partes 2, 3 y 4 conviene correrlas en tu computador (Whisper es pesado para la Pi).
"""

import os
import queue
import time
from datetime import datetime


DURATION_MINUTES = 80   # una clase dura 80 minutos
FS = 44100              # muestras por segundo (Hz)
CHANNELS = 1            # un solo micrófono

NOMBRE_BASE = datetime.now().strftime("clase_%Y-%m-%d_%H-%M")
ARCHIVO_AUDIO = f"{NOMBRE_BASE}.wav"

def grabar(nombre_archivo=ARCHIVO_AUDIO):
    """Graba DURATION_MINUTES minutos y va escribiendo en el disco
    mientras graba (así no se llena la memoria RAM)."""
    import sounddevice as sd
    import soundfile as sf

    cola = queue.Queue()

    def callback(indata, frames, tiempo, status):
        if status:
            print(status)
        cola.put(indata.copy())

    print(f"Grabando {DURATION_MINUTES} min en {nombre_archivo} ... (Ctrl+C para parar antes)")
    fin = time.time() + DURATION_MINUTES * 60

    with sf.SoundFile(nombre_archivo, mode="w", samplerate=FS, channels=CHANNELS) as archivo:
        with sd.InputStream(samplerate=FS, channels=CHANNELS, callback=callback):
            try:
                while time.time() < fin:
                    try:
                        archivo.write(cola.get(timeout=1))
                    except queue.Empty:
                        pass
            except KeyboardInterrupt:
                print("Grabación detenida manualmente.")

    print("Grabación terminada.")
    return nombre_archivo


#
def transcribir(ruta_audio):
    import whisper

    modelo = whisper.load_model("small")  # opciones: tiny, base, small, medium
    resultado = modelo.transcribe(ruta_audio, language="es")
    texto = resultado["text"]

    ruta_txt = os.path.splitext(ruta_audio)[0] + ".txt"
    with open(ruta_txt, "w", encoding="utf-8") as f:
        f.write(texto)
    print(f"Transcripción guardada en {ruta_txt}")
    return texto



# PARTE 3: resumir con Gemini (Google AI Studio)
# ---------------------------------------------------------------
# "gemini-flash-latest" siempre apunta al Flash más reciente.
# Si da error, mira los modelos disponibles con listar_modelos() y cambia este nombre.
MODELO_GEMINI = "gemini-flash-latest"


def listar_modelos():
    from google import genai

    cliente = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    for m in cliente.models.list():
        print(m.name)


def resumir(texto):
    from google import genai

    # La clave la sacas de https://aistudio.google.com/api-keys
    # NO la escribas directo en el código ni la subas a GitHub.
    cliente = genai.Client(api_key=os.environ["GEMINI_API_KEY"])

    prompt = (
        "Eres un asistente que toma apuntes de clases universitarias. "
        "Responde en español, claro y organizado.\n\n"
        "Esta es la transcripción de una clase. Hazme:\n"
        "1) Un resumen corto\n"
        "2) Los puntos clave\n"
        "3) Conceptos o definiciones importantes\n"
        "4) Tareas o fechas mencionadas (si las hay)\n\n"
        f"Transcripción:\n{texto}"
    )

    respuesta = cliente.models.generate_content(
        model=MODELO_GEMINI,
        contents=prompt,
    )
    return respuesta.text


# ---------------------------------------------------------------
# PARTE 4a: exportar a Word (lo más fácil)
# ---------------------------------------------------------------
def guardar_word(resumen, nombre=NOMBRE_BASE):
    from docx import Document

    doc = Document()
    doc.add_heading(f"Resumen de {nombre}", level=1)
    for linea in resumen.split("\n"):
        doc.add_paragraph(linea)
    ruta = f"{nombre}_resumen.docx"
    doc.save(ruta)
    print(f"Word guardado en {ruta}")
    return ruta


# PARTE 4b: exportar a Google Docs
# ---------------------------------------------------------------
SCOPES = ["https://www.googleapis.com/auth/documents"]


def servicio_docs():
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    creds = None
    # token_docs.json lo crea el programa solo, la primera vez que aceptas el permiso
    if os.path.exists("token_docs.json"):
        creds = Credentials.from_authorized_user_file("token_docs.json", SCOPES)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            # credentials.json es el que SÍ descargas de Google Cloud
            flow = InstalledAppFlow.from_client_secrets_file("credentials.json", SCOPES)
            creds = flow.run_local_server(port=0)  # abre el navegador
        with open("token_docs.json", "w") as f:
            f.write(creds.to_json())

    return build("docs", "v1", credentials=creds)


def guardar_google_docs(resumen, titulo=NOMBRE_BASE):
    docs = servicio_docs()
    doc = docs.documents().create(body={"title": f"Resumen {titulo}"}).execute()
    doc_id = doc["documentId"]

    docs.documents().batchUpdate(
        documentId=doc_id,
        body={"requests": [{"insertText": {"location": {"index": 1}, "text": resumen}}]},
    ).execute()

    url = f"https://docs.google.com/document/d/{doc_id}/edit"
    print(f"Google Doc creado: {url}")
    return url


# Programa principal
# ---------------------------------------------------------------
if __name__ == "__main__":
    # En la Raspberry Pi: solo grabar
    # grabar()

    # En tu computador (después de pasar el .wav con scp):
    # texto = transcribir("clase_2026-10-29_10-00.wav")
    # resumen = resumir(texto)
    # guardar_word(resumen)
    # guardar_google_docs(resumen)

    # Todo seguido (si lo corres todo en el mismo equipo):
    audio = grabar()
    texto = transcribir(audio)
    resumen = resumir(texto)
    guardar_word(resumen)
    guardar_google_docs(resumen)
